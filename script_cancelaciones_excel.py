import threading
import tkinter as tk
from tkinter import scrolledtext, messagebox
import os
import sys
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

# ─── FIX PARA EXE: apuntar a la carpeta estándar de navegadores ───────────────
# El .exe busca el navegador en su carpeta temporal interna; esto redirige a la
# carpeta estándar donde 'playwright install chromium' descarga los navegadores.
os.environ["PLAYWRIGHT_BROWSERS_PATH"] = str(
    Path.home() / "AppData" / "Local" / "ms-playwright"
)

from playwright.sync_api import sync_playwright

# ─── CONFIGURACIÓN ────────────────────────────────────────────────────────────
URL_LOGIN = "https://www.carvajaltys.mx/"
USUARIO   = "003321"
PASSWORD  = "Len0v02015"
# ──────────────────────────────────────────────────────────────────────────────


def obtener_folios(texto):
    """Parsea el texto pegado por el usuario.
    Formato esperado (una línea por folio):
    FOLIO,FECHA_INICIO,FECHA_FIN
    Ejemplo: 12345,01/01/2024,31/01/2024
    """
    folios = []
    for linea in texto.strip().split("\n"):
        partes = linea.strip().split(",")
        if len(partes) == 3:
            folios.append({
                "folio": partes[0].strip(),
                "desde": partes[1].strip(),
                "hasta": partes[2].strip()
            })
    return folios


def get_log_path():
    """Retorna la ruta del archivo de Excel junto al .exe o script."""
    if getattr(sys, 'frozen', False):
        carpeta = os.path.dirname(sys.executable)
    else:
        carpeta = os.path.dirname(os.path.abspath(__file__))
    fecha = datetime.now().strftime("%Y-%m-%d")
    return os.path.join(carpeta, f"cancelaciones_{fecha}.xlsx")


def log(widget, mensaje, mostrar=False):
    """Guarda el mensaje SOLO en Excel si es final (con ✅ o ⚠️).
    Solo lo muestra en la interfaz cuando mostrar=True (mensajes finales)."""
    if mostrar:
        widget.insert(tk.END, mensaje + "\n")
        widget.see(tk.END)
        # Guardar en Excel solo si es un mensaje final
        if "✅" in mensaje or "⚠️" in mensaje:
            guardar_en_excel(mensaje)


def guardar_en_excel(mensaje):
    """Añade el mensaje al archivo Excel del día con fecha y hora."""
    ruta = get_log_path()
    
    try:
        # Verificar si el archivo ya existe
        if os.path.exists(ruta):
            from openpyxl import load_workbook
            wb = load_workbook(ruta)
            ws = wb.active
        else:
            # Crear un nuevo workbook con encabezados
            wb = Workbook()
            ws = wb.active
            ws.title = "Cancelaciones"
            
            # Encabezados
            ws['A1'] = "Fecha"
            ws['B1'] = "Hora"
            ws['C1'] = "Resultado"
            
            # Estilos de encabezado
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            for cell in ['A1', 'B1', 'C1']:
                ws[cell].fill = header_fill
                ws[cell].font = header_font
                ws[cell].alignment = Alignment(horizontal="center", vertical="center")
        
        # Añadir fila con la fecha, hora y el mensaje
        fila = ws.max_row + 1
        fecha = datetime.now().strftime("%d/%m/%Y")
        hora = datetime.now().strftime("%H:%M:%S")
        ws[f'A{fila}'] = fecha
        ws[f'B{fila}'] = hora
        ws[f'C{fila}'] = mensaje
        ws[f'A{fila}'].alignment = Alignment(horizontal="center")
        ws[f'B{fila}'].alignment = Alignment(horizontal="center")
        ws[f'C{fila}'].alignment = Alignment(horizontal="left", wrap_text=True)
        
        # Ajustar ancho de columnas
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 80
        
        wb.save(ruta)
    except Exception as e:
        print(f"Error al guardar en Excel: {e}")


def automatizar(folios, log_widget, btn_iniciar):
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=False)
            page = browser.new_page()

            # ── PASO 1: Ir al login ───────────────────────────────────────
            log(log_widget, "🌐 Abriendo página de login...")
            page.goto(URL_LOGIN)
            page.wait_for_load_state("networkidle")

            # ── PASO 2: Seleccionar Monitor Emisor ────────────────────────
            log(log_widget, "🔽 Seleccionando Monitor Emisor...")
            page.select_option("#cmbSistemas", "7")
            page.wait_for_timeout(500)

            # ── PASO 3: Ingresar credenciales ─────────────────────────────
            log(log_widget, "🔐 Ingresando credenciales...")
            page.fill("#txtUseMonitorEmisor", USUARIO)
            page.fill("#txtPassMonitorEmisor", PASSWORD)
            page.click("#btnIngrear7")
            page.wait_for_load_state("networkidle")

            # ── PASO 4: Cerrar popup de aviso ─────────────────────────────
            try:
                page.click("td.close", timeout=5000)
                log(log_widget, "✅ Popup de aviso cerrado.")
                page.wait_for_timeout(800)
            except Exception:
                log(log_widget, "ℹ️ No apareció popup de aviso, continuando...")

            # ── PASO 5: Abrir panel de Búsqueda ──────────────────────────
            log(log_widget, "🔍 Abriendo panel de búsqueda...")
            page.click("#Busq")
            page.wait_for_timeout(1000)

            # ── PASO 6: Iterar por cada folio ────────────────────────────
            for i, item in enumerate(folios, 1):
                log(log_widget, f"\n📋 [{i}/{len(folios)}] Procesando folio: {item['folio']}")

                # Limpiar campos y llenar
                page.fill("#ctl00_ContentPlaceBody_txtFolio", "")
                page.fill("#ctl00_ContentPlaceBody_txtFolio", item["folio"])

                page.fill("#ctl00_ContentPlaceBody_TB_desde", "")
                page.fill("#ctl00_ContentPlaceBody_TB_desde", item["desde"])

                page.fill("#ctl00_ContentPlaceBody_TB_hasta", "")
                page.fill("#ctl00_ContentPlaceBody_TB_hasta", item["hasta"])

                # Buscar
                log(log_widget, "   🔎 Buscando...")
                page.click("#ctl00_ContentPlaceBody_btnBuscarOk")
                page.wait_for_load_state("networkidle")

                # Seleccionar SIEMPRE la casilla fija de la primera fila de resultados
                # (ctl02 = primera fila de datos de la tabla grvMonitor).
                # No se itera ni se adivina: si esta casilla no aparece, se salta el folio.
                try:
                    SEL_CASILLA = "#ctl00_ContentPlaceBody_grvMonitor_ctl02_checkbox"

                    log(log_widget, "   ⏳ Esperando resultados de la búsqueda...")
                    casilla = page.wait_for_selector(SEL_CASILLA, timeout=15000)

                    if not casilla.is_checked():
                        casilla.check(force=True)
                    log(log_widget, "   ✅ Casilla del resultado seleccionada (primera fila).")

                    # Esperar el postback que dispara el check (__doPostBack)
                    try:
                        page.wait_for_load_state("networkidle", timeout=8000)
                    except Exception:
                        pass
                    page.wait_for_timeout(1000)

                except Exception:
                    # Diagnóstico: ver qué quedó en la página tras la búsqueda
                    try:
                        n_grv = page.locator("[id*='grvMonitor']").count()
                        n_checks = page.locator("input[type='checkbox'][id*='grvMonitor']").count()
                        log(log_widget, f"   ⚠️ No apareció la casilla para el folio {item['folio']}, saltando...")
                        log(log_widget, f"   🔎 Elementos grvMonitor: {n_grv} | Casillas: {n_checks} | URL: {page.url}")
                    except Exception:
                        log(log_widget, f"   ⚠️ No apareció la casilla para el folio {item['folio']}, saltando...")
                    continue

                # Solicitar cancelación
                log(log_widget, "   🚫 Solicitando cancelación...")
                page.click("#ctl00_ContentPlaceBody_btnCancelacion")
                page.wait_for_load_state("networkidle")

                # ── PASO 7: Gestor de cancelaciones (se abre como VENTANA nueva) ──
                # El gestor NO es un iframe ni parte de la página principal: es una
                # ventana independiente (frmCancel.aspx). En Playwright aparece como
                # una nueva "page" dentro del mismo contexto del navegador.
                cancelacion_enviada = False
                try:
                    log(log_widget, "   📋 Esperando ventana del gestor de cancelaciones...")

                    def es_gestor(pg):
                        try:
                            if pg == page:
                                return False
                            if "frmCancel" in pg.url:
                                return True
                            return "Gestor de Cancelaciones" in (pg.title() or "")
                        except Exception:
                            return False

                    popup = None
                    for _ in range(40):  # hasta ~20 segundos
                        for pg_aux in page.context.pages:
                            if es_gestor(pg_aux):
                                popup = pg_aux
                                break
                        if popup:
                            break
                        page.wait_for_timeout(500)

                    if popup is None:
                        # Diagnóstico: listar todas las ventanas abiertas para
                        # ver qué se generó realmente al solicitar la cancelación
                        urls_abiertas = []
                        for pg_aux in page.context.pages:
                            try:
                                urls_abiertas.append(pg_aux.url)
                            except Exception:
                                pass
                        log(log_widget, "   ⚠️ No se detectó la ventana del gestor de cancelaciones.")
                        log(log_widget, f"   🔎 Ventanas abiertas en este momento: {urls_abiertas}")
                    else:
                        log(log_widget, f"   📋 Gestor abierto: {popup.url}")
                        popup.wait_for_load_state("domcontentloaded")
                        popup.wait_for_selector("#gvCancelaciones_ctl02_ddl_Motivo", timeout=15000)

                        # Seleccionar motivo: 02 - Comprobante emitido con errores sin relación
                        popup.select_option("#gvCancelaciones_ctl02_ddl_Motivo", "02")
                        log(log_widget, "   ✅ Motivo seleccionado: 02 - Sin relación.")

                        # El change dispara un __doPostBack. No esperamos networkidle/load,
                        # porque este sitio puede mantener la navegación pendiente aunque el
                        # botón ya esté disponible nuevamente.
                        popup.wait_for_timeout(2500)

                        # Esperar directamente al botón y enviarlo. La confirmación real es
                        # el texto de SweetAlert "Cancelación enviada", no el resultado de
                        # wait_for_load_state().
                        boton_enviar = popup.locator("#btnSendCancel")
                        boton_enviar.wait_for(state="visible", timeout=60000)
                        boton_enviar.click(timeout=10000, no_wait_after=True)

                        try:
                            popup.get_by_text("Cancelación", exact=False).wait_for(
                                state="visible", timeout=15000
                            )
                            texto_confirmacion = popup.locator("body").inner_text().lower()
                            if ("exitosa" in texto_confirmacion or
                                    "enviada" in texto_confirmacion or
                                    "enviado" in texto_confirmacion):
                                cancelacion_enviada = True
                                log(log_widget, "   ✅ Confirmación recibida: cancelación exitosa/enviada.")
                            else:
                                raise Exception("No se reconoció la confirmación de cancelación")
                        except Exception:
                            # Algunas respuestas cierran la ventana inmediatamente.
                            # Si el botón fue pulsado y la ventana se cerró, también se
                            # considera envío confirmado.
                            popup.wait_for_timeout(1500)
                            if popup.is_closed():
                                cancelacion_enviada = True
                                log(log_widget, "   ✅ Ventana cerrada después de enviar la cancelación.")
                            else:
                                raise

                        # La ventana muestra una alerta y suele cerrarse sola
                        popup.wait_for_timeout(2500)
                        try:
                            if popup.is_closed():
                                log(log_widget, "   ✔ Ventana del gestor se cerró automáticamente.")
                            else:
                                popup.close()
                                log(log_widget, "   ✔ Ventana del gestor cerrada.")
                        except Exception:
                            pass

                        # Devolver el foco a la página principal
                        try:
                            page.bring_to_front()
                        except Exception:
                            pass

                except Exception as e:
                    log(log_widget, f"   ⚠️ Error en gestor de cancelaciones: {str(e)}")

                # ── PASO 8: Volver a picarle Buscar para actualizar el estatus CFDI ────────────
                estado_cfdi = None
                try:
                    log(log_widget, "   🔄 Verificando estatus CFDI...")

                    def buscar_y_leer():
                        # Cada intento vuelve a ejecutar Buscar para forzar la
                        # actualización del Estado CFDI en el servidor.
                        page.locator("#ctl00_ContentPlaceBody_btnBuscarOk").click()
                        try:
                            page.wait_for_load_state("domcontentloaded", timeout=20000)
                        except Exception:
                            pass
                        # Esperar a que el postback reconstruya la tabla y el
                        # Estado CFDI vuelva a estar visible en la página.
                        try:
                            page.wait_for_selector(
                                "span[id*='lblEstadoCFDI']", state="visible", timeout=20000
                            )
                        except Exception:
                            pass
                        # Margen extra: el "Cargando" dura menos de un segundo y
                        # el postback puede terminar después de domcontentloaded.
                        page.wait_for_timeout(2000)

                        # Leer únicamente el Estado CFDI de la tabla ya actualizada.
                        return page.locator("span[id*='lblEstadoCFDI']").last.inner_text()

                    # Después de enviar la cancelación, el portal puede tardar en
                    # reflejar el cambio. Esperamos un minuto antes de consultar.
                    log(log_widget, "   ⏳ Esperando 60 segundos para que el portal actualice el estatus...")
                    page.wait_for_timeout(60000)

                    # Consultar nuevamente. Si aún no cambió, hacemos dos consultas
                    # adicionales separadas por 10 segundos, sin revisar fechas.
                    for intento in range(3):
                        estado_cfdi = buscar_y_leer()
                        log(log_widget, f"   🔎 Consulta {intento + 1}/3: Estatus CFDI = {estado_cfdi}")
                        if estado_cfdi and estado_cfdi.strip().lower() == "intermedio":
                            break
                        if intento < 2:
                            page.wait_for_timeout(10000)

                except Exception as e:
                    log(log_widget, f"   ⚠️ No se pudo verificar estatus para folio {item['folio']}: {str(e)}")

                # ── MENSAJE FINAL (lo único que el usuario ve en pantalla) ─────
                # La fuente de verdad es el Estado CFDI del portal: si ya dice
                # "Intermedio", la cancelación se confirmó aunque el popup de
                # confirmación no haya podido leerse.
                if estado_cfdi and estado_cfdi.strip().lower() == "intermedio":
                    log(log_widget, f"✅ FOLIO: {item['folio']} | Cancelación solicitada | Estatus CFDI: {estado_cfdi}", mostrar=True)
                elif not cancelacion_enviada:
                    log(log_widget, f"⚠️ FOLIO: {item['folio']} | No se pudo confirmar el envío de la cancelación.", mostrar=True)
                elif estado_cfdi:
                    log(log_widget, f"⚠️ FOLIO: {item['folio']} | Cancelación solicitada | Estatus CFDI aún: {estado_cfdi}", mostrar=True)
                else:
                    log(log_widget, f"✅ FOLIO: {item['folio']} | Cancelación solicitada | El portal aún no actualiza el estatus; revisar nuevamente en unos minutos.", mostrar=True)

                page.wait_for_timeout(500)

            browser.close()
            log(log_widget, "\n🎉 ¡Proceso terminado!")

    except Exception as e:
        log(log_widget, f"\n❌ Error: {str(e)}")
    finally:
        btn_iniciar.config(state="normal")


# ─── INTERFAZ GRÁFICA ─────────────────────────────────────────────────────────
def iniciar_ui():
    root = tk.Tk()
    root.title("Automatización - Cancelaciones CFDI")
    root.geometry("620x560")
    root.resizable(False, False)

    # Instrucciones
    tk.Label(
        root,
        text="Pega los folios en el formato:  FOLIO,FECHA_INICIO,FECHA_FIN",
        font=("Arial", 9),
        fg="gray"
    ).pack(pady=(10, 0))
    tk.Label(
        root,
        text="Ejemplo:  12345,01/01/2024,31/01/2024",
        font=("Arial", 9, "italic"),
        fg="gray"
    ).pack()

    # Área de entrada
    tk.Label(root, text="Folios:", font=("Arial", 10, "bold")).pack(anchor="w", padx=15, pady=(10, 0))
    entrada = scrolledtext.ScrolledText(root, height=10, font=("Courier", 10))
    entrada.pack(padx=15, fill=tk.BOTH)

    # Log
    tk.Label(root, text="Log:", font=("Arial", 10, "bold")).pack(anchor="w", padx=15, pady=(10, 0))
    log_widget = scrolledtext.ScrolledText(root, height=10, font=("Courier", 9), bg="#1e1e1e", fg="#00ff00")
    log_widget.pack(padx=15, pady=(0, 10), fill=tk.BOTH)

    # Botón iniciar
    def iniciar():
        texto = entrada.get("1.0", tk.END)
        folios = obtener_folios(texto)
        if not folios:
            messagebox.showwarning("Sin folios", "No se encontraron folios válidos.\nVerifica el formato: FOLIO,FECHA_INICIO,FECHA_FIN")
            return
        log_widget.delete("1.0", tk.END)
        log(log_widget, f"🚀 Iniciando con {len(folios)} folio(s)...\n")
        btn_iniciar.config(state="disabled")
        hilo = threading.Thread(target=automatizar, args=(folios, log_widget, btn_iniciar), daemon=True)
        hilo.start()

    btn_iniciar = tk.Button(
        root,
        text="▶  Iniciar Automatización",
        command=iniciar,
        bg="#28a745",
        fg="white",
        font=("Arial", 11, "bold"),
        height=2,
        cursor="hand2"
    )
    btn_iniciar.pack(padx=15, pady=(0, 15), fill=tk.X)

    root.mainloop()


if __name__ == "__main__":
    iniciar_ui()
